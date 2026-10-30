"""Browser input form <-> EComTool_User_Input.xlsx mapping.

Form values are written onto a copy of the shipped template and stored through
pipeline.save_user_input, so uploads and form saves share validation and cache clearing.
"""
from __future__ import annotations

import io
import warnings
from pathlib import Path
from typing import Any

import pipeline

INPUT_SHEET = "Input"
PARTICIPANT_FIRST_ROW = 35
PARTICIPANT_LAST_ROW = 84
PARTICIPANT_LIMIT = PARTICIPANT_LAST_ROW - PARTICIPANT_FIRST_ROW + 1
# Participant columns on the Input sheet; column A holds the fixed "Participant N" id.
PARTICIPANT_COLUMNS = {
    "category": "B", "type": "C", "number": "D",
    "annual_kwh": "E", "load_profile": "F", "comment": "G",
}

# Form key -> Input sheet cell. Single source of truth for read and write.
FIELD_CELLS = {
    "country": "A4", "city": "B4", "area_type": "C4",
    "tariff_from": "D4", "tariff_to": "E4",
    "pv_existing_kw": "B10", "pv_projected_kw": "B11", "pv_profile": "B12",
    "ev_existing": "B14", "ev_projected": "B15",
    "ev_home_profile": "B16", "ev_public_profile": "B17",
    "ev_workplace_profile": "B18", "ev_fast_profile": "B19",
    "ev_home_pct": "B20", "ev_public_pct": "B21",
    "ev_workplace_pct": "B22", "ev_fast_pct": "B23",
    "smart_charging_pct": "B24",
    "smart_ev_home_profile": "B25", "smart_ev_public_profile": "B26",
    "bess_existing_kwh": "B28", "bess_projected_kwh": "B29",
    "tariff_from_own": "B91", "tariff_to_own": "B92",
    "pv_age_years": "B93", "pv_degradation_pct": "B94", "bess_degradation_pct": "B95",
    "horizon_years": "B96", "pv_trajectory": "B97", "ev_trajectory": "B98",
    "load_growth_pct": "B99",
}


def template_path() -> Path:
    """Shipped template, by absolute path: memory Stage 1 changes the process cwd."""
    return (pipeline.STAGE1_DIR / pipeline.USER_INPUT_NAME).resolve()


def _load(path: Path, *, data_only: bool):
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
        return load_workbook(path, data_only=data_only)


def _text(value: Any) -> Any:
    return value.strip() if isinstance(value, str) else value


def read_form(path: Path | None = None) -> dict[str, Any]:
    """Return form values from the active input workbook (or `path`)."""
    workbook = _load(path or pipeline.user_input_path(), data_only=True)
    try:
        sheet = workbook[INPUT_SHEET]
        form: dict[str, Any] = {key: _text(sheet[cell].value) for key, cell in FIELD_CELLS.items()}
        participants = []
        for row in range(PARTICIPANT_FIRST_ROW, PARTICIPANT_LAST_ROW + 1):
            item = {key: _text(sheet[f"{col}{row}"].value) for key, col in PARTICIPANT_COLUMNS.items()}
            if item["type"] not in (None, ""):  # active row = Type set, not the prefilled id
                participants.append(item)
        form["participants"] = participants
        return form
    finally:
        workbook.close()


def read_options() -> dict[str, Any]:
    """Dropdown lists from the shipped template's helper sheets."""
    workbook = _load(template_path(), data_only=True)
    try:
        def columns(name: str) -> dict[str, list[str]]:
            result: dict[str, list[str]] = {}
            for column in workbook[name].iter_cols(values_only=True):
                head, *items = column
                if head not in (None, ""):
                    result[str(head).strip()] = [str(v).strip() for v in items if v not in (None, "")]
            return result

        return {
            "cities": columns("Geography"),
            "participant_types": columns("Consumers"),
            "profiles": columns("Profiles"),
            "tariffs": columns("Tariffs"),
        }
    finally:
        workbook.close()


class FormError(ValueError):
    """Form rejected; `errors` lists every problem, one per line in str()."""

    def __init__(self, errors: list[str]):
        super().__init__("\n".join(errors))
        self.errors = errors


AREA_TYPES = ("Urban", "Suburban", "Rural")
HORIZON_YEARS = (5, 10, 12, 15, 18, 20)
TRAJECTORIES = ("Instantaneous", "Linear", "Exponential")
GREEN_TARIFF = "Green Tariff"  # Stage 1 supports it for Ukrainian export (To Grid) only
DATABASE_PROFILE = "Database"
OWN_PROFILE_HINT = "Own profiles are not supported in the form; upload an input file instead"
EV_SHARE_KEYS = ("ev_home_pct", "ev_public_pct", "ev_workplace_pct", "ev_fast_pct")
# Profile fields the form only accepts as plain "Database" (PV is checked against its two Database options).
DATABASE_ONLY_PROFILES = (
    "ev_home_profile", "ev_public_profile", "ev_workplace_profile", "ev_fast_profile",
    "smart_ev_home_profile", "smart_ev_public_profile",
)


def _label(key: str) -> str:
    return f"{key} (Input!{FIELD_CELLS[key]})"


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and value == value \
        and value not in (float("inf"), float("-inf"))


def _check_number(errors: list[str], name: str, value: Any, *, low: float | None = None,
                  high: float | None = None, integer: bool = False, positive: bool = False,
                  optional: bool = False) -> None:
    if value is None or value == "":
        if not optional:
            errors.append(f"{name}: required")
        return
    if not _is_number(value):
        errors.append(f"{name}: must be a finite number")
        return
    if integer and value != int(value):
        errors.append(f"{name}: must be a whole number")
    if positive and value <= 0:
        errors.append(f"{name}: must be greater than 0")
    if low is not None and value < low:
        errors.append(f"{name}: must be at least {low:g}")
    if high is not None and value > high:
        errors.append(f"{name}: must be at most {high:g}")


def _check_choice(errors: list[str], name: str, value: Any, allowed) -> bool:
    if value in allowed:
        return True
    errors.append(f"{name}: '{value}' is not one of {', '.join(map(str, allowed))}")
    return False


def validate_form(form: dict[str, Any], options: dict[str, Any] | None = None) -> list[str]:
    """Return every problem found; empty list means the form can be saved.

    Bounds follow the Excel template where Stage 1 does not enforce its own; Stage 1 rules
    (B92 > 0, Green Tariff Ukraine export only, Type-driven participant rows) are mirrored.
    """
    options = options or read_options()
    errors: list[str] = []
    missing = [key for key in FIELD_CELLS if key not in form]
    if "participants" not in form:
        missing.append("participants")
    if missing:
        return [f"Missing form fields: {', '.join(missing)}"]

    country = form["country"]
    tariffs = options["tariffs"].get(country, [])
    if _check_choice(errors, _label("country"), country, list(options["cities"])):
        _check_choice(errors, _label("city"), form["city"], options["cities"][country])
        _check_choice(errors, _label("tariff_from"), form["tariff_from"],
                      [t for t in tariffs if t != GREEN_TARIFF])
        _check_choice(errors, _label("tariff_to"), form["tariff_to"], tariffs)
    _check_choice(errors, _label("area_type"), form["area_type"], AREA_TYPES)

    pv_options = [p for p in options["profiles"]["PV Profile"] if p.startswith(DATABASE_PROFILE)]
    for key, allowed in [("pv_profile", pv_options)] + [(k, [DATABASE_PROFILE]) for k in DATABASE_ONLY_PROFILES]:
        value = form[key]
        if isinstance(value, str) and value.startswith("Own"):
            errors.append(f"{_label(key)}: {OWN_PROFILE_HINT}")
        else:
            _check_choice(errors, _label(key), value, allowed)

    for key in ("pv_existing_kw", "pv_projected_kw", "bess_existing_kwh", "bess_projected_kwh"):
        _check_number(errors, _label(key), form[key], low=0)
    for key in ("ev_existing", "ev_projected"):
        _check_number(errors, _label(key), form[key], low=0, integer=True)
    for key in EV_SHARE_KEYS + ("smart_charging_pct",):
        _check_number(errors, _label(key), form[key], low=0, high=100)
    shares = [form[key] for key in EV_SHARE_KEYS]
    if all(_is_number(v) for v in shares) and abs(sum(shares) - 100) > 1e-9:
        errors.append("EV charging shares (Input!B20:B23): must sum to 100")

    _check_number(errors, _label("tariff_from_own"), form["tariff_from_own"], low=0, optional=True)
    _check_number(errors, _label("tariff_to_own"), form["tariff_to_own"], positive=True)  # Stage 1 needs it always
    _check_number(errors, _label("pv_age_years"), form["pv_age_years"], low=0, high=25)
    _check_number(errors, _label("pv_degradation_pct"), form["pv_degradation_pct"], low=0.1, high=2)
    _check_number(errors, _label("bess_degradation_pct"), form["bess_degradation_pct"], low=0.1, high=2)
    _check_number(errors, _label("load_growth_pct"), form["load_growth_pct"], low=0.01, high=100, optional=True)
    _check_choice(errors, _label("horizon_years"), form["horizon_years"], HORIZON_YEARS)
    _check_choice(errors, _label("pv_trajectory"), form["pv_trajectory"], TRAJECTORIES)
    _check_choice(errors, _label("ev_trajectory"), form["ev_trajectory"], TRAJECTORIES)

    participants = form["participants"]
    if not isinstance(participants, list) or not participants:
        errors.append("participants: add at least one participant")
        return errors
    if len(participants) > PARTICIPANT_LIMIT:
        errors.append(f"participants: at most {PARTICIPANT_LIMIT} are supported")
    for index, item in enumerate(participants[:PARTICIPANT_LIMIT], start=1):
        name = f"Participant {index}"
        category, kind = item.get("category"), item.get("type")
        if _check_choice(errors, f"{name} category", category, list(options["participant_types"])):
            _check_choice(errors, f"{name} type", kind, options["participant_types"][category])
        _check_number(errors, f"{name} number", item.get("number"), low=1, integer=True, optional=True)
        _check_number(errors, f"{name} annual consumption, kWh", item.get("annual_kwh"), low=0, optional=True)
        profile = item.get("load_profile")
        if isinstance(profile, str) and profile.startswith("Own"):
            errors.append(f"{name} load profile: {OWN_PROFILE_HINT}")
        else:
            _check_choice(errors, f"{name} load profile", profile, [DATABASE_PROFILE])
        comment = item.get("comment")
        if comment is not None and not isinstance(comment, str):
            errors.append(f"{name} comment: must be text")
    return errors


def build_workbook(form: dict[str, Any]) -> bytes:
    """Write form onto a fresh copy of the shipped template; return XLSX bytes."""
    unknown = set(form) - set(FIELD_CELLS) - {"participants"}
    if unknown:
        raise ValueError(f"Unknown form fields: {', '.join(sorted(unknown))}")
    participants = form.get("participants")
    if participants is not None and len(participants) > PARTICIPANT_LIMIT:
        raise ValueError(f"At most {PARTICIPANT_LIMIT} participants are supported")

    workbook = _load(template_path(), data_only=False)  # formulas kept
    try:
        sheet = workbook[INPUT_SHEET]
        for key, value in form.items():
            if key in FIELD_CELLS:
                sheet[FIELD_CELLS[key]].value = value
        if participants is not None:
            # Replace the whole range so template sample rows cannot survive a deletion.
            for index, row in enumerate(range(PARTICIPANT_FIRST_ROW, PARTICIPANT_LAST_ROW + 1)):
                sheet[f"A{row}"].value = f"Participant {index + 1}"
                for col in PARTICIPANT_COLUMNS.values():
                    sheet[f"{col}{row}"].value = None
            for index, item in enumerate(participants):
                row = PARTICIPANT_FIRST_ROW + index
                for key, col in PARTICIPANT_COLUMNS.items():
                    value = item.get(key)
                    sheet[f"{col}{row}"].value = None if value == "" else value
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()


def save_form(form: dict[str, Any]) -> Path:
    """Validate, build the workbook and store it as the active input."""
    errors = validate_form(form)
    if errors:
        raise FormError(errors)
    return pipeline.save_user_input(build_workbook(form))
